"""Build the Sabil Tracker overview deck.

Design language mirrors the app's light theme: near-white canvas, white cards with
hairline borders, violet used sparingly as an accent, emerald/rose for money semantics.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- design tokens

INK = RGBColor(0x1A, 0x1A, 0x24)
MUTED = RGBColor(0x6B, 0x6B, 0x7E)
FAINT = RGBColor(0x8E, 0x8E, 0x9E)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
VIOLET_MID = RGBColor(0x8B, 0x5C, 0xF6)
VIOLET_TINT = RGBColor(0xF5, 0xF3, 0xFF)
VIOLET_LINE = RGBColor(0xDD, 0xD6, 0xFE)
CANVAS = RGBColor(0xFA, 0xFA, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HAIRLINE = RGBColor(0xE7, 0xE5, 0xF0)
ROSE = RGBColor(0xBE, 0x12, 0x3C)
ROSE_TINT = RGBColor(0xFF, 0xF1, 0xF3)
ROSE_LINE = RGBColor(0xFB, 0xCF, 0xD8)
EMERALD = RGBColor(0x04, 0x78, 0x57)
EMERALD_TINT = RGBColor(0xEC, 0xFD, 0xF5)
EMERALD_LINE = RGBColor(0xA7, 0xF3, 0xD0)

FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_W = SLIDE_W - 2 * MARGIN

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- primitives


def new_slide(canvas=CANVAS):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = canvas
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def text(
    slide,
    x,
    y,
    w,
    h,
    content,
    size=15,
    color=MUTED,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line=1.25,
    space_after=0,
):
    """content: str, or list of (text, {overrides}) tuples rendered as paragraphs."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    items = [(content, {})] if isinstance(content, str) else content

    for i, (line_text, over) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = over.get("align", align)
        p.line_spacing = over.get("line", line)
        p.space_after = Pt(over.get("space_after", space_after))
        run = p.add_run()
        run.text = line_text
        f = run.font
        f.name = FONT
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.color.rgb = over.get("color", color)
    return box


def card(slide, x, y, w, h, fill=WHITE, border=HAIRLINE, radius=0.055):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.adjustments[0] = radius
    shp.text_frame.text = ""
    return shp


def pill(slide, x, y, w, h, label, fill=VIOLET_TINT, border=VIOLET_LINE, color=VIOLET, size=11):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.adjustments[0] = 0.5
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return shp


def header(slide, eyebrow, title, sub=None):
    """Standard slide header: violet eyebrow, big title, optional subtitle."""
    y = Inches(0.62)
    text(slide, MARGIN, y, CONTENT_W, Inches(0.25), eyebrow.upper(), size=11.5, color=VIOLET, bold=True)
    text(slide, MARGIN, y + Inches(0.32), CONTENT_W, Inches(0.6), title, size=30, color=INK, bold=True, line=1.05)
    if sub:
        text(slide, MARGIN, y + Inches(1.0), Inches(9.6), Inches(0.5), sub, size=14.5, color=MUTED, line=1.3)
        return y + Inches(1.62)
    return y + Inches(1.18)


def footer(slide, number):
    text(slide, MARGIN, SLIDE_H - Inches(0.55), Inches(4), Inches(0.25), "Sabil Tracker", size=10, color=FAINT)
    text(
        slide,
        SLIDE_W - MARGIN - Inches(1),
        SLIDE_H - Inches(0.55),
        Inches(1),
        Inches(0.25),
        str(number),
        size=10,
        color=FAINT,
        align=PP_ALIGN.RIGHT,
    )


def icon_card(slide, x, y, w, h, emoji, title, body, accent=VIOLET):
    card(slide, x, y, w, h)
    text(slide, x + Inches(0.28), y + Inches(0.26), Inches(0.6), Inches(0.42), emoji, size=21, color=INK)
    text(slide, x + Inches(0.28), y + Inches(0.78), w - Inches(0.56), Inches(0.3), title, size=14.5, color=INK, bold=True)
    text(slide, x + Inches(0.28), y + Inches(1.13), w - Inches(0.56), h - Inches(1.3), body, size=11.5, color=MUTED, line=1.28)


# ---------------------------------------------------------------- slide 1: title

s = new_slide(WHITE)

band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.11), SLIDE_H)
band.fill.solid()
band.fill.fore_color.rgb = VIOLET
band.line.fill.background()
band.shadow.inherit = False

blob = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.6), Inches(-1.5), Inches(6.6), Inches(6.6))
blob.fill.solid()
blob.fill.fore_color.rgb = VIOLET_TINT
blob.line.fill.background()
blob.shadow.inherit = False

pill(s, MARGIN, Inches(1.55), Inches(3.5), Inches(0.34), "PERSONAL FINANCE + ASSETS")

text(s, MARGIN, Inches(2.15), Inches(9.4), Inches(1.9),
     "Sabil Tracker", size=58, color=INK, bold=True, line=1.0)
text(s, MARGIN, Inches(3.25), Inches(8.9), Inches(1.5),
     "Track cash, bank, assets and documents\nin one place — and never miss a deadline.",
     size=21, color=MUTED, line=1.3)

rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(4.62), Inches(1.5), Pt(3))
rule.fill.solid()
rule.fill.fore_color.rgb = VIOLET
rule.line.fill.background()
rule.shadow.inherit = False

text(s, MARGIN, Inches(5.0), Inches(9.4), Inches(1.0),
     [("Why the old way fails  ·  What this does  ·  Why it's worth using", {"size": 14, "color": FAINT})])

# ---------------------------------------------------------------- slide 2: problem intro

s = new_slide()
header(s, "The problem", "You're not disorganised.\nYour tools just weren't built for this.")

stmts = [
    ("Money lives in two places.", "Cash in your pocket, the rest in the bank. Neither view is the real one."),
    ("Deadlines live nowhere.", "Passports, licences, insurance and servicing expire on their own schedule."),
    ("Details have no home.", "An odometer reading or a document number doesn't fit in \"amount + note\"."),
]
x = MARGIN
w = Inches(3.65)
gap = Inches(0.31)
for i, (t, b) in enumerate(stmts):
    cx = x + i * (w + gap)
    card(s, cx, Inches(2.55), w, Inches(2.35))
    num = pill(s, cx + Inches(0.28), Inches(2.8), Inches(0.42), Inches(0.42), str(i + 1))
    text(s, cx + Inches(0.28), Inches(3.45), w - Inches(0.56), Inches(0.5), t, size=16, color=INK, bold=True, line=1.15)
    text(s, cx + Inches(0.28), Inches(4.05), w - Inches(0.56), Inches(0.8), b, size=12, color=MUTED, line=1.3)

text(s, MARGIN, Inches(5.35), CONTENT_W, Inches(0.6),
     "The result is the same every time: you find out too late, and it costs money.",
     size=15.5, color=INK, bold=True)
footer(s, 2)

# ---------------------------------------------------------------- slide 3: the six pains

s = new_slide()
header(s, "The old way", "Six problems this replaces")

pains = [
    ("💸", "Balance guesswork", "Cash counted in your head, bank checked in another app. The real number is a guess."),
    ("📉", "The spreadsheet that dies", "Accurate for two weeks, then it drifts, falls out of sync and gets abandoned."),
    ("🛂", "Silent expiries", "Passports, licences and insurance expire quietly — you find out at the airport."),
    ("🚗", "Maintenance by memory", "Servicing and tyre changes tracked in your head, so they slip until something breaks."),
    ("🧩", "Nowhere for the details", "Generic apps store an amount and a note. No odometer, no document number, no expiry."),
    ("🔥", "Overspending, found late", "No guardrails, so the overspend is only visible after the month has closed."),
]
w = Inches(3.65)
h = Inches(2.0)
gapx = Inches(0.31)
gapy = Inches(0.26)
for i, (emoji, t, b) in enumerate(pains):
    cx = MARGIN + (i % 3) * (w + gapx)
    cy = Inches(2.35) + (i // 3) * (h + gapy)
    icon_card(s, cx, cy, w, h, emoji, t, b)
footer(s, 3)

# ---------------------------------------------------------------- slide 4: comparison

s = new_slide()
header(s, "Before / after", "The same six problems, solved")

col_w = Inches(5.55)
left_x = MARGIN
right_x = MARGIN + col_w + Inches(0.35)
head_y = Inches(2.2)

pill(s, left_x, head_y, Inches(1.7), Inches(0.34), "THE OLD WAY", ROSE_TINT, ROSE_LINE, ROSE)
pill(s, right_x, head_y, Inches(2.15), Inches(0.34), "WITH SABIL TRACKER", EMERALD_TINT, EMERALD_LINE, EMERALD)

rows = [
    ("Two balances, mentally reconciled", "One ledger; cash and bank reconcile on save"),
    ("A spreadsheet you stop updating", "A logger built for daily use — seconds per entry"),
    ("Expiry dates you remember, or don't", "Every date tracked, colour-coded by urgency"),
    ("Service intervals tracked by feel", "Odometer + service dates with next-due badges"),
    ("Fields you can't add without a developer", "Define your own fields per category. No code"),
    ("Overspend discovered after the fact", "Live budget progress and over-budget alerts"),
]
y = Inches(2.75)
rh = Inches(0.62)
for i, (old, fresh) in enumerate(rows):
    cy = y + i * (rh + Inches(0.11))
    card(s, left_x, cy, col_w, rh, ROSE_TINT, ROSE_LINE)
    text(s, left_x + Inches(0.24), cy + Inches(0.16), col_w - Inches(0.48), rh, old, size=12, color=ROSE, line=1.15)
    card(s, right_x, cy, col_w, rh, EMERALD_TINT, EMERALD_LINE)
    text(s, right_x + Inches(0.24), cy + Inches(0.16), col_w - Inches(0.48), rh, fresh, size=12, color=EMERALD, line=1.15)
footer(s, 4)

# ---------------------------------------------------------------- slide 5: what it is

s = new_slide()
header(s, "What it is", "One place for money, assets and documents",
       "A self-hosted tracker that treats a passport, a car service and a grocery run as things worth recording properly.")

pillars = [
    ("💰", "A real ledger", "Physical cash and bank accounts tracked separately, reconciled automatically on every entry."),
    ("🗂️", "A flexible schema", "Categories you nest, and custom fields you define per category — without touching code."),
    ("🔔", "A deadline engine", "Every date you record can raise a reminder, on your own schedule, before it matters."),
]
w = Inches(3.65)
for i, (emoji, t, b) in enumerate(pillars):
    cx = MARGIN + i * (w + Inches(0.31))
    icon_card(s, cx, Inches(3.1), w, Inches(2.15), emoji, t, b)
footer(s, 5)

# ---------------------------------------------------------------- slide 6: feature map

s = new_slide()
header(s, "Features", "What's in the application")

feats = [
    ("💵", "Dual cash / bank ledger", "Separate balances, reconciled in real time."),
    ("🌳", "Hierarchical categories", "Parent → sub-category, each with its own emoji."),
    ("🧩", "Dynamic Fields Engine", "Define Date / Number / Text / Checkbox fields."),
    ("🛂", "Expiry & maintenance planner", "Days-remaining, colour-coded by urgency."),
    ("🔔", "Automated reminders", "Nightly scheduler, email + in-app alerts."),
    ("🎯", "Budgets", "Live progress bars and over-budget warnings."),
    ("📊", "Reports & CSV export", "Drill-down charts, filters, downloadable statements."),
    ("🔒", "Self-hosted", "Your database, your server, your data."),
]
w = Inches(3.03)
h = Inches(1.72)
for i, (emoji, t, b) in enumerate(feats):
    cx = MARGIN + (i % 4) * (w + Inches(0.2))
    cy = Inches(2.3) + (i // 4) * (h + Inches(0.26))
    card(s, cx, cy, w, h)
    text(s, cx + Inches(0.24), cy + Inches(0.24), Inches(0.6), Inches(0.4), emoji, size=19, color=INK)
    text(s, cx + Inches(0.24), cy + Inches(0.72), w - Inches(0.48), Inches(0.5), t, size=13, color=INK, bold=True, line=1.12)
    text(s, cx + Inches(0.24), cy + Inches(1.18), w - Inches(0.48), Inches(0.5), b, size=10.5, color=MUTED, line=1.25)
footer(s, 6)

# ---------------------------------------------------------------- slide 7: dual ledger

s = new_slide()
top = header(s, "Feature 01", "Cash and bank, finally reconciled",
             "Most tools treat all money as one pool. Physical cash is where budgets quietly break.")

card(s, MARGIN, Inches(2.95), Inches(5.55), Inches(2.72))
text(s, MARGIN + Inches(0.35), Inches(3.2), Inches(4.8), Inches(0.4), "How it works", size=13, color=VIOLET, bold=True)
text(s, MARGIN + Inches(0.35), Inches(3.65), Inches(4.85), Inches(1.7),
     [("Every account is typed as cash, bank or credit.", {"space_after": 7}),
      ("Logging an expense decrements that account; income increments it.", {"space_after": 7}),
      ("Editing or deleting reverses the old effect and reapplies the new one.", {"space_after": 7}),
      ("Balance maths runs inside a database transaction, so it can't half-apply.", {})],
     size=12.5, color=MUTED, line=1.25)

rx = MARGIN + Inches(5.9)
card(s, rx, Inches(2.95), Inches(5.55), Inches(2.72))
text(s, rx + Inches(0.35), Inches(3.2), Inches(4.8), Inches(0.4), "What you see", size=13, color=VIOLET, bold=True)
tiles = [("💵  Cash in hand", EMERALD), ("🏦  Bank & digital", VIOLET), ("💰  Total net wealth", INK)]
for i, (label, col) in enumerate(tiles):
    ty = Inches(3.68) + i * Inches(0.55)
    card(s, rx + Inches(0.35), ty, Inches(4.85), Inches(0.45), CANVAS, HAIRLINE)
    text(s, rx + Inches(0.55), ty + Inches(0.11), Inches(3.2), Inches(0.3), label, size=12, color=col, bold=True)
    text(s, rx + Inches(3.5), ty + Inches(0.11), Inches(1.55), Inches(0.3), "live", size=11, color=FAINT, align=PP_ALIGN.RIGHT)

text(s, MARGIN, Inches(5.95), CONTENT_W, Inches(0.5),
     "Why it matters: you can answer \"what can I actually spend today?\" without doing arithmetic.",
     size=14, color=INK, bold=True)
footer(s, 7)

# ---------------------------------------------------------------- slide 8: dynamic fields

s = new_slide()
header(s, "Feature 02  ·  the differentiator", "Define the fields each thing actually needs",
       "A passport needs a number and an expiry. A car needs an odometer and a service date. No app ships with your fields — so you define them.")

card(s, MARGIN, Inches(3.05), Inches(6.1), Inches(2.75))
text(s, MARGIN + Inches(0.35), Inches(3.3), Inches(5.4), Inches(0.4),
     "Add a field to any sub-category", size=13, color=VIOLET, bold=True)
rows = [
    ("Odometer (km)", "Number"),
    ("Last Service Date", "Date  ·  reminder on"),
    ("Tyre Brand", "Text"),
    ("Inspection done?", "Checkbox"),
]
for i, (name, kind) in enumerate(rows):
    ry = Inches(3.8) + i * Inches(0.47)
    card(s, MARGIN + Inches(0.35), ry, Inches(5.4), Inches(0.38), CANVAS, HAIRLINE)
    text(s, MARGIN + Inches(0.55), ry + Inches(0.08), Inches(2.9), Inches(0.28), name, size=11.5, color=INK, bold=True)
    text(s, MARGIN + Inches(3.5), ry + Inches(0.08), Inches(2.1), Inches(0.28), kind, size=11, color=VIOLET, align=PP_ALIGN.RIGHT)

rx = MARGIN + Inches(6.45)
card(s, rx, Inches(3.05), Inches(5.0), Inches(2.75), VIOLET_TINT, VIOLET_LINE)
text(s, rx + Inches(0.35), Inches(3.3), Inches(4.3), Inches(0.4), "What that buys you", size=13, color=VIOLET, bold=True)
text(s, rx + Inches(0.35), Inches(3.8), Inches(4.3), Inches(1.9),
     [("The logger renders those fields automatically when you pick the category.", {"space_after": 8}),
      ("Any Date field can raise a reminder, with your own days-before offset.", {"space_after": 8}),
      ("Reports grow columns to match the fields you filtered on.", {"space_after": 8}),
      ("No code. No migration. No developer.", {"bold": True, "color": VIOLET})],
     size=12, color=MUTED, line=1.25)
footer(s, 8)

# ---------------------------------------------------------------- slide 9: planner + reminders

s = new_slide()
header(s, "Feature 03", "Nothing expires without warning",
       "Every date you record becomes a tracked deadline, surfaced long before it bites.")

cards = [
    ("🟢", "More than 180 days", "Green. Nothing to do yet.", EMERALD, EMERALD_TINT, EMERALD_LINE),
    ("🟠", "30 – 180 days", "Amber. Start the renewal.", RGBColor(0xB4, 0x53, 0x09), RGBColor(0xFF, 0xF7, 0xED), RGBColor(0xFE, 0xD7, 0xAA)),
    ("🔴", "Under 30 days", "Red, and it pulses. Act now.", ROSE, ROSE_TINT, ROSE_LINE),
]
w = Inches(3.65)
for i, (emoji, t, b, col, fill, line_c) in enumerate(cards):
    cx = MARGIN + i * (w + Inches(0.31))
    card(s, cx, Inches(3.05), w, Inches(1.5), fill, line_c)
    text(s, cx + Inches(0.28), Inches(3.28), Inches(0.5), Inches(0.35), emoji, size=16, color=INK)
    text(s, cx + Inches(0.85), Inches(3.3), w - Inches(1.1), Inches(0.3), t, size=14, color=col, bold=True)
    text(s, cx + Inches(0.28), Inches(3.82), w - Inches(0.56), Inches(0.5), b, size=11.5, color=MUTED)

card(s, MARGIN, Inches(4.8), CONTENT_W, Inches(1.35))
text(s, MARGIN + Inches(0.35), Inches(5.02), Inches(3.2), Inches(0.35), "How the alert reaches you", size=13, color=VIOLET, bold=True)
steps = [
    "A nightly scheduled command scans every active reminder",
    "Anything inside its days-before window is queued",
    "You get an email and an in-app alert; the dashboard shows it first",
]
for i, stp in enumerate(steps):
    sx = MARGIN + Inches(0.35) + i * Inches(3.85)
    text(s, sx, Inches(5.45), Inches(3.6), Inches(0.6),
         [("→  " if i else "", {}), (stp, {})], size=11.5, color=MUTED, line=1.25)
footer(s, 9)

# ---------------------------------------------------------------- slide 10: budgets + reports

s = new_slide()
header(s, "Feature 04", "Guardrails while it still matters")

card(s, MARGIN, Inches(2.35), Inches(5.55), Inches(3.1))
text(s, MARGIN + Inches(0.35), Inches(2.6), Inches(4.8), Inches(0.4), "Budgets", size=15, color=INK, bold=True)
text(s, MARGIN + Inches(0.35), Inches(3.05), Inches(4.85), Inches(0.5),
     "Set a monthly limit per category. Progress updates as you log.", size=12, color=MUTED, line=1.25)
bars = [("Under 70%", "on track", EMERALD, 0.55), ("70 – 90%", "watch it", RGBColor(0xB4, 0x53, 0x09), 0.8), ("Over 100%", "alert", ROSE, 1.0)]
for i, (label, note, col, frac) in enumerate(bars):
    by = Inches(3.72) + i * Inches(0.52)
    text(s, MARGIN + Inches(0.35), by, Inches(1.6), Inches(0.25), label, size=11, color=INK, bold=True)
    text(s, MARGIN + Inches(4.05), by, Inches(1.15), Inches(0.25), note, size=10.5, color=col, align=PP_ALIGN.RIGHT)
    track = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN + Inches(0.35), by + Inches(0.26), Inches(4.85), Inches(0.14))
    track.fill.solid(); track.fill.fore_color.rgb = RGBColor(0xEE, 0xEC, 0xF5)
    track.line.fill.background(); track.shadow.inherit = False; track.adjustments[0] = 0.5
    fillbar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN + Inches(0.35), by + Inches(0.26), Inches(4.85 * frac), Inches(0.14))
    fillbar.fill.solid(); fillbar.fill.fore_color.rgb = col
    fillbar.line.fill.background(); fillbar.shadow.inherit = False; fillbar.adjustments[0] = 0.5

rx = MARGIN + Inches(5.9)
card(s, rx, Inches(2.35), Inches(5.55), Inches(3.1))
text(s, rx + Inches(0.35), Inches(2.6), Inches(4.8), Inches(0.4), "Reports & statements", size=15, color=INK, bold=True)
text(s, rx + Inches(0.35), Inches(3.05), Inches(4.85), Inches(2.2),
     [("Filter by date range and any set of categories.", {"space_after": 9}),
      ("Expense-by-category donut with drill-down from parent into sub-category.", {"space_after": 9}),
      ("A trend line that separates everyday spending from asset maintenance cost.", {"space_after": 9}),
      ("Tables that grow columns to show your custom fields.", {"space_after": 9}),
      ("One-click CSV export of the filtered statement.", {})],
     size=12, color=MUTED, line=1.25)
footer(s, 10)

# ---------------------------------------------------------------- slide 11: how it works

s = new_slide()
header(s, "How it works", "Three steps, then it runs itself")

steps = [
    ("01", "Set it up once", "Add your cash and bank accounts, then build the category tree you actually use — and give the categories that need them their own custom fields."),
    ("02", "Log as you go", "Pick expense, income or asset log. Choose the category, and the fields that matter for that thing appear automatically."),
    ("03", "Get warned early", "Balances reconcile instantly. Budgets track live. Deadlines surface on the dashboard and by email before they land."),
]
w = Inches(3.65)
for i, (num, t, b) in enumerate(steps):
    cx = MARGIN + i * (w + Inches(0.31))
    card(s, cx, Inches(2.5), w, Inches(2.75))
    text(s, cx + Inches(0.3), Inches(2.78), Inches(1.2), Inches(0.6), num, size=30, color=VIOLET_LINE, bold=True)
    text(s, cx + Inches(0.3), Inches(3.45), w - Inches(0.6), Inches(0.4), t, size=16, color=INK, bold=True)
    text(s, cx + Inches(0.3), Inches(3.92), w - Inches(0.6), Inches(1.2), b, size=12, color=MUTED, line=1.3)
footer(s, 11)

# ---------------------------------------------------------------- slide 12: why use it

s = new_slide()
header(s, "Why use it", "What actually changes for you")

outcomes = [
    ("You always know what you can spend today.", "Not what the bank app says — what you have, across cash and bank, right now."),
    ("Nothing expires without warning.", "Renewals start weeks early instead of the day you're turned away at a counter."),
    ("Your tracker matches your life, not a template.", "If it matters to you, there's a field for it — and you added it yourself."),
    ("Your data stays yours.", "Self-hosted. Your database, your server, no vendor pricing or outages in the way."),
]
w = Inches(5.55)
h = Inches(1.42)
for i, (t, b) in enumerate(outcomes):
    cx = MARGIN + (i % 2) * (w + Inches(0.35))
    cy = Inches(2.4) + (i // 2) * (h + Inches(0.28))
    card(s, cx, cy, w, h)
    tick = pill(s, cx + Inches(0.3), cy + Inches(0.3), Inches(0.34), Inches(0.34), "✓")
    text(s, cx + Inches(0.82), cy + Inches(0.32), w - Inches(1.1), Inches(0.35), t, size=14.5, color=INK, bold=True, line=1.1)
    text(s, cx + Inches(0.82), cy + Inches(0.78), w - Inches(1.1), Inches(0.5), b, size=11.5, color=MUTED, line=1.25)
footer(s, 12)

# ---------------------------------------------------------------- slide 13: architecture

s = new_slide()
header(s, "Under the hood", "Built on a boring, proven stack",
       "A single Laravel application serving a React SPA — no separate API to keep in sync.")

layers = [
    ("Browser", "React 19 · TypeScript · Tailwind CSS 4", "Charts hand-built as inline SVG — no chart library."),
    ("Bridge", "Inertia.js v3", "Server-side routing with an SPA feel. No REST layer, no CORS."),
    ("Backend", "Laravel 13 · PHP 8.4", "Thin controllers, form requests, service classes, scheduled commands."),
    ("Storage", "SQLite (Postgres / MariaDB ready)", "Five tables, soft deletes, JSON columns for the custom fields."),
]
y = Inches(2.9)
for i, (label, tech, note) in enumerate(layers):
    cy = y + i * Inches(0.82)
    card(s, MARGIN, cy, CONTENT_W, Inches(0.68))
    pill(s, MARGIN + Inches(0.25), cy + Inches(0.16), Inches(1.15), Inches(0.35), label)
    text(s, MARGIN + Inches(1.65), cy + Inches(0.2), Inches(4.0), Inches(0.3), tech, size=12.5, color=INK, bold=True)
    text(s, MARGIN + Inches(5.9), cy + Inches(0.21), Inches(5.5), Inches(0.3), note, size=11.5, color=MUTED)
footer(s, 13)

# ---------------------------------------------------------------- slide 14: status

s = new_slide()
header(s, "Where it stands", "Built and working today")

stats = [
    ("9", "application screens", "Dashboard, transactions, categories,\naccounts, planner, reminders,\nbudgets, reports, landing."),
    ("65", "automated tests passing", "Covering balances, cascade deletes,\nreminders, budgets, reports\nand CSV export."),
    ("5", "core data tables", "Accounts, categories, transactions,\nreminders and budgets —\nall user-scoped."),
]
w = Inches(3.65)
for i, (n, label, note) in enumerate(stats):
    cx = MARGIN + i * (w + Inches(0.31))
    card(s, cx, Inches(2.45), w, Inches(2.45))
    text(s, cx + Inches(0.3), Inches(2.7), w - Inches(0.6), Inches(0.8), n, size=44, color=VIOLET, bold=True, line=1.0)
    text(s, cx + Inches(0.3), Inches(3.55), w - Inches(0.6), Inches(0.3), label, size=13, color=INK, bold=True)
    text(s, cx + Inches(0.3), Inches(3.95), w - Inches(0.6), Inches(0.9), note, size=11, color=MUTED, line=1.3)

card(s, MARGIN, Inches(5.15), CONTENT_W, Inches(0.85), VIOLET_TINT, VIOLET_LINE)
text(s, MARGIN + Inches(0.35), Inches(5.38), CONTENT_W - Inches(0.7), Inches(0.45),
     "Seeded demo data ships with the app — a passport expiring in 40 days, a tyre change due in 5, six months of transactions — so the reminder states are visible on first run.",
     size=12.5, color=VIOLET, line=1.25)
footer(s, 14)

# ---------------------------------------------------------------- slide 15: close

s = new_slide(WHITE)

blob = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2.2), Inches(3.4), Inches(7.5), Inches(7.5))
blob.fill.solid(); blob.fill.fore_color.rgb = VIOLET_TINT
blob.line.fill.background(); blob.shadow.inherit = False

band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, SLIDE_W - Inches(0.11), 0, Inches(0.11), SLIDE_H)
band.fill.solid(); band.fill.fore_color.rgb = VIOLET
band.line.fill.background(); band.shadow.inherit = False

text(s, MARGIN, Inches(2.5), Inches(10.5), Inches(1.6),
     "Stop guessing your balance.\nStop finding out too late.",
     size=40, color=INK, bold=True, line=1.15)
text(s, MARGIN, Inches(4.35), Inches(9.0), Inches(0.9),
     "One ledger for cash and bank. Your own fields on the things you track.\nReminders that fire before the deadline, not after.",
     size=16, color=MUTED, line=1.35)

rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(5.5), Inches(1.5), Pt(3))
rule.fill.solid(); rule.fill.fore_color.rgb = VIOLET
rule.line.fill.background(); rule.shadow.inherit = False

text(s, MARGIN, Inches(5.85), Inches(8), Inches(0.4), "Sabil Tracker", size=15, color=VIOLET, bold=True)

out = "/Users/jasim/P/sabil/tracker/expense-tracker/docs/sabil-tracker-overview.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
